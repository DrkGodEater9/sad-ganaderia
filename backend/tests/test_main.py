"""Pruebas del backend. No dependen de credenciales de Earth Engine: los
endpoints que llaman a earth-engine/modelo como subproceso se prueban solo
hasta donde llega la validacion (geojson, EE_PROJECT), sin invocar
Earth Engine de verdad."""

import json

import pytest
from fastapi.testclient import TestClient

import main


@pytest.fixture
def cliente(tmp_path, monkeypatch):
    geojson = {
        "type": "FeatureCollection",
        "features": [
            {"type": "Feature", "properties": {"nombre": "Potrero 1"}, "geometry": {"type": "Polygon", "coordinates": [[[0, 0], [0, 1], [1, 1], [1, 0], [0, 0]]]}},
            {"type": "Feature", "properties": {"nombre": "Potrero 2"}, "geometry": {"type": "Polygon", "coordinates": [[[2, 0], [2, 1], [3, 1], [3, 0], [2, 0]]]}},
        ],
    }
    ruta_geojson = tmp_path / "potreros.geojson"
    ruta_geojson.write_text(json.dumps(geojson), encoding="utf-8")

    data_dir = tmp_path / "data"
    monkeypatch.setattr(main, "GEOJSON_PATH", ruta_geojson)
    monkeypatch.setattr(main, "DATA_DIR", data_dir)
    monkeypatch.setattr(main, "ESTADO_PATH", data_dir / "potreros_estado.xlsx")
    monkeypatch.setattr(main, "AFORO_PATH", data_dir / "aforo_campo.xlsx")
    monkeypatch.setattr(main, "INDICES_HISTORIAL_PATH", data_dir / "indices_historial.xlsx")
    monkeypatch.setattr(main, "INDICES_TEMP_PATH", data_dir / "_indices_temp.xlsx")
    monkeypatch.setattr(main, "MODELO_JOBLIB_PATH", data_dir / "modelo_entrenado.joblib")
    monkeypatch.setattr(main, "FINCA_NOMBRE", "Finca de prueba")
    monkeypatch.delenv("EE_PROJECT", raising=False)

    return TestClient(main.app)


def test_get_potreros_siembra_estado_inicial(cliente):
    respuesta = cliente.get("/api/potreros")
    assert respuesta.status_code == 200
    datos = respuesta.json()
    assert datos["finca"] == "Finca de prueba"
    assert datos["actualizado"] is None
    assert datos["modelo"] is None
    nombres = {p["nombre"] for p in datos["potreros"]}
    assert nombres == {"Potrero 1", "Potrero 2"}
    assert all(p["estado"] == "sin_datos" for p in datos["potreros"])


def test_guardar_geojson_reescribe_geojson_y_reinicia_estado(cliente):
    # Primero deja el estado con datos "viejos" (potrero 1 con estado verde),
    # como si ya hubiera corrido una prediccion antes.
    main.guardar_estado(
        {
            "finca": "Finca de prueba",
            "actualizado": "2026-09-01T00:00:00",
            "potreros": [
                {"nombre": "Potrero 1", "estado": "verde", "biomasa_kg_ha": 600, "fecha": "2026-09-01", "fuente_prediccion": "modelo_entrenado"}
            ],
        }
    )

    nuevo_geojson = {
        "type": "FeatureCollection",
        "features": [
            {"type": "Feature", "properties": {"nombre": "Potrero A"}, "geometry": {"type": "Polygon", "coordinates": [[[0, 0], [0, 1], [1, 1], [1, 0], [0, 0]]]}},
            {"type": "Feature", "properties": {"nombre": "Potrero B"}, "geometry": {"type": "Polygon", "coordinates": [[[2, 0], [2, 1], [3, 1], [3, 0], [2, 0]]]}},
        ],
    }
    respuesta = cliente.post("/api/geojson", json=nuevo_geojson)
    assert respuesta.status_code == 200
    datos = respuesta.json()

    assert datos["finca"] == "Finca de prueba"  # se conserva
    nombres = {p["nombre"] for p in datos["potreros"]}
    assert nombres == {"Potrero A", "Potrero B"}
    assert all(p["estado"] == "sin_datos" for p in datos["potreros"])  # se reinicia, no arrastra el viejo

    # El GeoJSON en disco tambien quedo actualizado.
    guardado = json.loads(main.GEOJSON_PATH.read_text(encoding="utf-8"))
    assert {f["properties"]["nombre"] for f in guardado["features"]} == {"Potrero A", "Potrero B"}

    # Y GET /api/potreros ya refleja los potreros nuevos.
    respuesta_get = cliente.get("/api/potreros")
    assert {p["nombre"] for p in respuesta_get.json()["potreros"]} == {"Potrero A", "Potrero B"}


def test_guardar_geojson_sin_potreros_con_nombre_falla_claro(cliente):
    respuesta = cliente.post(
        "/api/geojson",
        json={"type": "FeatureCollection", "features": [{"type": "Feature", "properties": {}, "geometry": None}]},
    )
    assert respuesta.status_code == 400


def test_eliminar_potrero_lo_saca_del_geojson_y_del_estado(cliente):
    cliente.get("/api/potreros")  # crea el estado inicial (Potrero 1 y 2, sin_datos)

    respuesta = cliente.delete("/api/potreros/Potrero 1")
    assert respuesta.status_code == 200
    datos = respuesta.json()
    assert {p["nombre"] for p in datos["potreros"]} == {"Potrero 2"}

    geojson_guardado = json.loads(main.GEOJSON_PATH.read_text(encoding="utf-8"))
    assert {f["properties"]["nombre"] for f in geojson_guardado["features"]} == {"Potrero 2"}

    respuesta_get = cliente.get("/api/potreros")
    assert {p["nombre"] for p in respuesta_get.json()["potreros"]} == {"Potrero 2"}


def test_eliminar_potrero_inexistente_da_404(cliente):
    respuesta = cliente.delete("/api/potreros/No Existe")
    assert respuesta.status_code == 404


def test_eliminar_potrero_preserva_estadisticas_del_modelo(cliente):
    from openpyxl import load_workbook

    main.cargar_estado()  # crea Potreros/Finca

    libro = load_workbook(main.ESTADO_PATH)
    hoja_modelo = libro.create_sheet("Modelo_Info")
    hoja_modelo.append(["algoritmo", "rmse"])
    hoja_modelo.append(["random_forest", 64.5])
    libro.save(main.ESTADO_PATH)

    cliente.delete("/api/potreros/Potrero 1")

    respuesta = cliente.get("/api/potreros")
    assert respuesta.json()["modelo"]["algoritmo"] == "random_forest"


def test_registrar_aforo_valido(cliente):
    respuesta = cliente.post(
        "/api/potreros/Potrero 1/aforo",
        json={"altura_cm_1": 12.5, "altura_cm_2": 14.0, "altura_cm_3": 13.2},
    )
    assert respuesta.status_code == 200
    cuerpo = respuesta.json()
    assert cuerpo["potrero"] == "Potrero 1"
    assert cuerpo["fecha"]  # se autocompleta con hoy

    filas = main.leer_hoja_como_dicts(main.AFORO_PATH, "Aforo")
    assert len(filas) == 1
    assert filas[0]["potrero"] == "Potrero 1"
    assert filas[0]["altura_cm_1"] == 12.5


@pytest.mark.parametrize("campo,valor", [("altura_cm_1", 0), ("altura_cm_1", 250), ("altura_cm_2", -5)])
def test_registrar_aforo_altura_fuera_de_rango(cliente, campo, valor):
    cuerpo = {"altura_cm_1": 12.5, "altura_cm_2": 14.0, "altura_cm_3": 13.2}
    cuerpo[campo] = valor
    respuesta = cliente.post("/api/potreros/Potrero 1/aforo", json=cuerpo)
    assert respuesta.status_code == 400
    assert "cm" in respuesta.json()["detail"]


def test_registrar_aforo_potrero_inexistente(cliente):
    respuesta = cliente.post(
        "/api/potreros/Potrero Fantasma/aforo",
        json={"altura_cm_1": 10, "altura_cm_2": 10, "altura_cm_3": 10},
    )
    assert respuesta.status_code == 404


def test_registrar_aforo_fecha_invalida(cliente):
    respuesta = cliente.post(
        "/api/potreros/Potrero 1/aforo",
        json={"altura_cm_1": 10, "altura_cm_2": 10, "altura_cm_3": 10, "fecha": "21-09-2026"},
    )
    assert respuesta.status_code == 400
    assert "AAAA-MM-DD" in respuesta.json()["detail"]


def test_registrar_aforo_dos_mediciones_se_acumulan(cliente):
    for _ in range(2):
        cliente.post(
            "/api/potreros/Potrero 1/aforo",
            json={"altura_cm_1": 10, "altura_cm_2": 10, "altura_cm_3": 10},
        )
    assert len(main.leer_hoja_como_dicts(main.AFORO_PATH, "Aforo")) == 2


def test_predecir_sin_geojson(cliente, monkeypatch, tmp_path):
    monkeypatch.setattr(main, "GEOJSON_PATH", tmp_path / "no_existe.geojson")
    respuesta = cliente.post("/api/predecir")
    assert respuesta.status_code == 400
    assert "GeoJSON" in respuesta.json()["detail"]


def test_predecir_sin_ee_project(cliente):
    respuesta = cliente.post("/api/predecir")
    assert respuesta.status_code == 400
    assert "EE_PROJECT" in respuesta.json()["detail"]


def test_guardar_y_cargar_estado_ida_y_vuelta(cliente):
    estado = {
        "finca": "Finca de prueba",
        "actualizado": "2026-09-21T10:00:00",
        "potreros": [
            {"nombre": "Potrero 1", "estado": "verde", "biomasa_kg_ha": 620.4, "fecha": "2026-09-18", "fuente_prediccion": "modelo_entrenado"},
            {"nombre": "Potrero 2", "estado": "rojo", "biomasa_kg_ha": 210.0, "fecha": "2026-09-18", "fuente_prediccion": "formula_provisional"},
        ],
    }
    main.guardar_estado(estado)
    recargado = main.cargar_estado()

    assert recargado["finca"] == "Finca de prueba"
    assert recargado["actualizado"] == "2026-09-21T10:00:00"
    potrero_1 = next(p for p in recargado["potreros"] if p["nombre"] == "Potrero 1")
    assert potrero_1["estado"] == "verde"
    assert potrero_1["biomasa_kg_ha"] == 620.4
    assert potrero_1["fuente_prediccion"] == "modelo_entrenado"


def test_cargar_estado_tolera_edicion_manual_del_excel(cliente):
    from openpyxl import load_workbook

    main.cargar_estado()  # crea el archivo inicial
    libro = load_workbook(main.ESTADO_PATH)
    hoja = libro["Potreros"]
    hoja["B2"] = "verde"
    hoja["C2"] = 999
    libro.save(main.ESTADO_PATH)

    recargado = main.cargar_estado()
    assert recargado["potreros"][0]["estado"] == "verde"
    assert recargado["potreros"][0]["biomasa_kg_ha"] == 999


def test_get_potreros_incluye_estadisticas_del_modelo(cliente):
    from openpyxl import Workbook

    main.cargar_estado()  # crea Potreros/Finca

    libro = Workbook()
    libro.active.title = "Potreros"
    libro.active.append(["nombre", "estado", "biomasa_kg_ha", "fecha", "fuente_prediccion"])
    libro.active.append(["Potrero 1", "verde", 620.4, "2026-09-18", "modelo_entrenado"])
    libro.active.append(["Potrero 2", "rojo", 210.0, "2026-09-18", "modelo_entrenado"])

    hoja_finca = libro.create_sheet("Finca")
    hoja_finca.append(["finca", "actualizado"])
    hoja_finca.append(["Finca de prueba", "2026-09-21T10:00:00"])

    hoja_modelo = libro.create_sheet("Modelo_Info")
    hoja_modelo.append(["algoritmo", "rmse", "r2", "n_muestras_entrenamiento", "rmse_random_forest", "rmse_regresion_lineal"])
    hoja_modelo.append(["random_forest", 64.5, 0.81, 14, 64.5, 148.9])

    hoja_importancia = libro.create_sheet("Modelo_Importancia")
    hoja_importancia.append(["variable", "importancia"])
    hoja_importancia.append(["NDVI_mean", 0.62])
    hoja_importancia.append(["SAVI_mean", 0.38])

    libro.save(main.ESTADO_PATH)

    respuesta = cliente.get("/api/potreros")
    datos = respuesta.json()

    assert datos["modelo"]["algoritmo"] == "random_forest"
    assert datos["modelo"]["rmse_random_forest"] == 64.5
    assert datos["modelo"]["rmse_regresion_lineal"] == 148.9
    assert datos["modelo"]["importancias"] == [
        {"variable": "NDVI_mean", "importancia": 0.62},
        {"variable": "SAVI_mean", "importancia": 0.38},
    ]


def test_estado_invalido_en_excel_cae_a_sin_datos(cliente):
    from openpyxl import load_workbook

    main.cargar_estado()
    libro = load_workbook(main.ESTADO_PATH)
    libro["Potreros"]["B2"] = "un_valor_cualquiera"
    libro.save(main.ESTADO_PATH)

    recargado = main.cargar_estado()
    assert recargado["potreros"][0]["estado"] == "sin_datos"
