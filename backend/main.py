"""Backend minimo (FastAPI) que orquesta las piezas ya existentes del
proyecto: el GeoJSON de potreros, el script de indices de Earth Engine
(earth-engine/extraer_indices.py) y el script de modelo de prediccion de
biomasa (modelo/entrenar_predecir.py).

No reimplementa la logica de esos scripts: los invoca como subprocesos y
lee/escribe los mismos archivos de datos que ellos usan.
"""

import json
import os
import subprocess
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel

BASE_DIR = Path(__file__).resolve().parent
REPO_DIR = BASE_DIR.parent
DATA_DIR = BASE_DIR / "data"

# Toda la informacion de trabajo se guarda en Excel dentro del proyecto
# (nada de base de datos): se puede abrir, revisar y hasta corregir a mano
# con Excel. El backend lee las hojas por nombre de columna, no por
# posicion, para tolerar ese tipo de ediciones manuales.
ESTADO_PATH = DATA_DIR / "potreros_estado.xlsx"
AFORO_PATH = DATA_DIR / "aforo_campo.xlsx"

# indices_historial.xlsx es un archivo que solo CRECE: cada corrida de
# /api/predecir le agrega una fila por potrero (con la fecha en que se
# consulto y la fecha real de la escena de satelite usada). El modelo
# necesita ese historial para poder cruzar cada medicion de aforo con el
# indice de satelite de esa MISMA fecha y potrero. _indices_temp.xlsx es
# solo un archivo de paso: ahi escribe earth-engine/extraer_indices.py en
# cada corrida, y de inmediato se copia (agregandose) al historial.
INDICES_HISTORIAL_PATH = DATA_DIR / "indices_historial.xlsx"
INDICES_TEMP_PATH = DATA_DIR / "_indices_temp.xlsx"

MODELO_JOBLIB_PATH = DATA_DIR / "modelo_entrenado.joblib"

EARTH_ENGINE_SCRIPT = REPO_DIR / "earth-engine" / "extraer_indices.py"
MODELO_SCRIPT = REPO_DIR / "modelo" / "entrenar_predecir.py"

ESTADOS_VALIDOS = {"verde", "ambar", "rojo", "sin_datos"}
FUENTES_PREDICCION_VALIDAS = {"modelo_entrenado", "formula_provisional"}


def cargar_variables_de_entorno_locales():
    """Lee backend/.env (si existe) sin depender de python-dotenv."""
    ruta_env = BASE_DIR / ".env"
    if not ruta_env.exists():
        return
    for linea in ruta_env.read_text(encoding="utf-8").splitlines():
        linea = linea.strip()
        if not linea or linea.startswith("#") or "=" not in linea:
            continue
        clave, _, valor = linea.partition("=")
        os.environ.setdefault(clave.strip(), valor.strip().strip('"').strip("'"))


cargar_variables_de_entorno_locales()

GEOJSON_PATH = Path(os.environ.get("GEOJSON_PATH", str(REPO_DIR / "potreros.geojson")))
FINCA_NOMBRE = os.environ.get("FINCA_NOMBRE", "Mi finca")
CORS_ORIGINS = os.environ.get("CORS_ORIGINS", "*")

# Archivo de clima diario (formato NASA POWER: columnas YEAR, DOY, T2M,
# PRECTOTCORR) opcional. Sin el, el modelo entrena solo con los indices
# de satelite (sin variables climaticas).
_clima_env = os.environ.get("CLIMA_PATH", "").strip()
CLIMA_PATH = Path(_clima_env) if _clima_env else None

# Umbrales del semaforo (kg de materia seca por hectarea), configurables
# sin tocar codigo -- son un punto de partida razonable, no una ley fija.
UMBRAL_ROJO = float(os.environ.get("UMBRAL_ROJO", "300"))
UMBRAL_AMBAR = float(os.environ.get("UMBRAL_AMBAR", "500"))

app = FastAPI(title="SAD - backend")

origenes = ["*"] if CORS_ORIGINS.strip() == "*" else [o.strip() for o in CORS_ORIGINS.split(",")]
app.add_middleware(
    CORSMiddleware,
    allow_origins=origenes,
    allow_methods=["*"],
    allow_headers=["*"],
    allow_credentials=False,
)


class AforoIn(BaseModel):
    altura_cm_1: float
    altura_cm_2: float
    altura_cm_3: float
    fecha: Optional[str] = None


class GeoJSONIn(BaseModel):
    type: str
    features: list[dict]


def cargar_geojson() -> dict:
    if not GEOJSON_PATH.exists():
        raise HTTPException(
            400,
            "No se encontro el GeoJSON de potreros en "
            f"'{GEOJSON_PATH}'. Exportalo desde SAD y colocalo "
            "ahi, o configura la variable de entorno GEOJSON_PATH con la "
            "ruta correcta.",
        )
    try:
        return json.loads(GEOJSON_PATH.read_text(encoding="utf-8"))
    except json.JSONDecodeError as error:
        raise HTTPException(400, f"El archivo GeoJSON de potreros no es un JSON valido: {error}")


def nombres_de_potreros() -> list[str]:
    geojson = cargar_geojson()
    nombres = []
    for feature in geojson.get("features", []):
        nombre = (feature.get("properties") or {}).get("nombre")
        if nombre:
            nombres.append(nombre)
    return nombres


def leer_hoja_como_dicts(ruta: Path, nombre_hoja: str) -> list[dict]:
    """Lee una hoja de Excel usando la primera fila como encabezados, para
    que el orden/posicion de las columnas no importe (tolera ediciones a
    mano en Excel)."""
    libro = load_workbook(ruta, data_only=True)
    if nombre_hoja not in libro.sheetnames:
        return []
    filas = list(libro[nombre_hoja].iter_rows(values_only=True))
    if not filas:
        return []
    encabezados = [str(c).strip() if c is not None else "" for c in filas[0]]
    resultado = []
    for fila in filas[1:]:
        if all(valor is None for valor in fila):
            continue
        resultado.append({encabezados[i]: fila[i] for i in range(min(len(encabezados), len(fila)))})
    return resultado


def normalizar_fecha(valor) -> Optional[str]:
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    return str(valor)


def cargar_estado() -> dict:
    if not ESTADO_PATH.exists():
        # Primera vez: no hay estado guardado todavia. Se arma uno inicial
        # a partir del GeoJSON, con cada potrero en "sin_datos" hasta que
        # se corra una prediccion.
        estado = {
            "finca": FINCA_NOMBRE,
            "actualizado": None,
            "modelo": None,
            "potreros": [
                {
                    "nombre": nombre,
                    "estado": "sin_datos",
                    "biomasa_kg_ha": None,
                    "fecha": None,
                    "fuente_prediccion": None,
                }
                for nombre in nombres_de_potreros()
            ],
        }
        guardar_estado(estado)
        return estado

    filas_finca = leer_hoja_como_dicts(ESTADO_PATH, "Finca")
    finca = (filas_finca[0].get("finca") if filas_finca else None) or FINCA_NOMBRE
    actualizado = normalizar_fecha(filas_finca[0].get("actualizado")) if filas_finca else None

    potreros = []
    for fila in leer_hoja_como_dicts(ESTADO_PATH, "Potreros"):
        nombre = fila.get("nombre")
        if not nombre:
            continue
        estado_valor = str(fila.get("estado") or "").strip().lower()
        if estado_valor not in ESTADOS_VALIDOS:
            estado_valor = "sin_datos"
        biomasa = fila.get("biomasa_kg_ha")
        biomasa = float(biomasa) if isinstance(biomasa, (int, float)) else None
        fuente = str(fila.get("fuente_prediccion") or "").strip().lower() or None
        if fuente not in FUENTES_PREDICCION_VALIDAS:
            fuente = None
        potreros.append(
            {
                "nombre": str(nombre),
                "estado": estado_valor,
                "biomasa_kg_ha": biomasa,
                "fecha": normalizar_fecha(fila.get("fecha")),
                "fuente_prediccion": fuente,
            }
        )

    # Hoja opcional "Modelo_Info": la escribe el script de modelo con
    # metadatos de transparencia (que algoritmo gano, que tan bueno es,
    # cuando se entreno). Si no existe todavia (o el modelo esta usando
    # la formula provisional), simplemente no hay nada que mostrar aqui.
    filas_modelo = leer_hoja_como_dicts(ESTADO_PATH, "Modelo_Info")
    modelo_info = filas_modelo[0] if filas_modelo else None
    if modelo_info:
        modelo_info = {str(k): v for k, v in modelo_info.items()}
        # Hoja opcional "Modelo_Importancia": que variables pesaron mas en
        # el algoritmo que gano (vacia mientras se use la formula provisional).
        modelo_info["importancias"] = [
            {"variable": fila.get("variable"), "importancia": fila.get("importancia")}
            for fila in leer_hoja_como_dicts(ESTADO_PATH, "Modelo_Importancia")
            if fila.get("variable")
        ]

    return {"finca": finca, "actualizado": actualizado, "modelo": modelo_info, "potreros": potreros}


def guardar_estado(estado: dict):
    """Escribe potreros_estado.xlsx. Preserva las hojas Modelo_Info /
    Modelo_Importancia si ya existian (las escribe modelo/entrenar_predecir.py,
    no esta funcion) -- borrar un potrero o subir un geojson nuevo no debe
    tirar a la basura las estadisticas del ultimo modelo entrenado."""
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    hojas_a_preservar = ("Modelo_Info", "Modelo_Importancia")
    filas_preservadas = {}
    if ESTADO_PATH.exists():
        libro_anterior = load_workbook(ESTADO_PATH, data_only=True)
        for nombre_hoja in hojas_a_preservar:
            if nombre_hoja in libro_anterior.sheetnames:
                filas_preservadas[nombre_hoja] = list(libro_anterior[nombre_hoja].iter_rows(values_only=True))

    libro = Workbook()

    hoja_potreros = libro.active
    hoja_potreros.title = "Potreros"
    hoja_potreros.append(["nombre", "estado", "biomasa_kg_ha", "fecha", "fuente_prediccion"])
    for potrero in estado["potreros"]:
        hoja_potreros.append(
            [
                potrero["nombre"],
                potrero["estado"],
                potrero.get("biomasa_kg_ha"),
                potrero.get("fecha"),
                potrero.get("fuente_prediccion"),
            ]
        )

    hoja_finca = libro.create_sheet("Finca")
    hoja_finca.append(["finca", "actualizado"])
    hoja_finca.append([estado.get("finca", FINCA_NOMBRE), estado.get("actualizado")])

    for nombre_hoja, filas in filas_preservadas.items():
        hoja = libro.create_sheet(nombre_hoja)
        for fila in filas:
            hoja.append(list(fila))

    libro.save(ESTADO_PATH)


def guardar_fila_aforo(potrero: str, fecha: str, h1: float, h2: float, h3: float):
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if AFORO_PATH.exists():
        libro = load_workbook(AFORO_PATH)
        hoja = libro["Aforo"] if "Aforo" in libro.sheetnames else libro.active
    else:
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Aforo"
        hoja.append(["potrero", "fecha", "altura_cm_1", "altura_cm_2", "altura_cm_3"])

    hoja.append([potrero, fecha, h1, h2, h3])
    libro.save(AFORO_PATH)


def agregar_indices_al_historial(ruta_temp: Path, fecha_consulta: str):
    """Copia las filas que acaba de escribir earth-engine/extraer_indices.py
    (un archivo temporal, se sobreescribe cada corrida) al historial
    acumulado que usa el modelo para entrenar (indices_historial.xlsx,
    que solo crece)."""
    libro_temp = load_workbook(ruta_temp, data_only=True)
    hoja_temp = libro_temp.active
    filas = list(hoja_temp.iter_rows(values_only=True))
    if not filas:
        return
    encabezados = [str(c).strip() if c is not None else "" for c in filas[0]]

    if INDICES_HISTORIAL_PATH.exists():
        libro_historial = load_workbook(INDICES_HISTORIAL_PATH)
        hoja_historial = libro_historial.active
    else:
        libro_historial = Workbook()
        hoja_historial = libro_historial.active
        hoja_historial.title = "Indices"
        hoja_historial.append(encabezados + ["fecha_consulta"])

    for fila in filas[1:]:
        if all(valor is None for valor in fila):
            continue
        hoja_historial.append(list(fila) + [fecha_consulta])

    libro_historial.save(INDICES_HISTORIAL_PATH)
    ruta_temp.unlink(missing_ok=True)


def ultimas_lineas(texto: str, cantidad: int = 8) -> str:
    lineas = [l for l in texto.strip().splitlines() if l.strip()]
    return "\n".join(lineas[-cantidad:]) if lineas else "No se recibio ningun detalle del proceso."


@app.get("/api/potreros")
def obtener_potreros():
    return cargar_estado()


@app.post("/api/geojson")
def guardar_geojson(body: GeoJSONIn):
    """Guarda el GeoJSON dibujado en la pestana 'Dibujar potreros' como el
    que usa el resto de la app (earth-engine, modelo, /api/predecir) --
    sin que el usuario tenga que descargar el archivo y volver a subirlo a
    mano. Reinicia el estado de los potreros a 'sin_datos' porque cambio la
    geometria: los indices satelitales viejos ya no aplican necesariamente."""
    nombres = [
        (feature.get("properties") or {}).get("nombre")
        for feature in body.features
        if (feature.get("properties") or {}).get("nombre")
    ]
    if not nombres:
        raise HTTPException(400, "El GeoJSON no tiene ningun potrero con nombre. Dibuja y cierra al menos uno.")

    GEOJSON_PATH.parent.mkdir(parents=True, exist_ok=True)
    GEOJSON_PATH.write_text(json.dumps(body.model_dump(), ensure_ascii=False, indent=2), encoding="utf-8")

    finca_actual = cargar_estado().get("finca", FINCA_NOMBRE) if ESTADO_PATH.exists() else FINCA_NOMBRE
    estado = {
        "finca": finca_actual,
        "actualizado": None,
        "modelo": None,
        "potreros": [
            {"nombre": nombre, "estado": "sin_datos", "biomasa_kg_ha": None, "fecha": None, "fuente_prediccion": None}
            for nombre in nombres
        ],
    }
    guardar_estado(estado)
    return estado


@app.delete("/api/potreros/{nombre}")
def eliminar_potrero(nombre: str):
    """Borra un potrero de todos lados: del GeoJSON (para que no vuelva a
    aparecer en la proxima prediccion) y del estado guardado (para que
    desaparezca de inmediato de 'Mis potreros', 'Registrar medicion' y
    'Estadisticas'). El aforo y los indices historicos de ese potrero NO se
    borran -- son datos de campo reales, quedan como historial aunque el
    potrero ya no exista."""
    geojson = cargar_geojson()
    features_restantes = [
        f for f in geojson.get("features", []) if (f.get("properties") or {}).get("nombre") != nombre
    ]
    if len(features_restantes) == len(geojson.get("features", [])):
        raise HTTPException(404, f"No existe un potrero llamado '{nombre}' en el GeoJSON de la finca.")

    geojson["features"] = features_restantes
    GEOJSON_PATH.write_text(json.dumps(geojson, ensure_ascii=False, indent=2), encoding="utf-8")

    estado = cargar_estado()
    estado["potreros"] = [p for p in estado["potreros"] if p["nombre"] != nombre]
    guardar_estado(estado)
    return cargar_estado()


@app.post("/api/potreros/{nombre}/aforo")
def registrar_aforo(nombre: str, body: AforoIn):
    nombres_validos = {p["nombre"] for p in cargar_estado()["potreros"]}
    if nombres_validos and nombre not in nombres_validos:
        raise HTTPException(404, f"No existe un potrero llamado '{nombre}' en el GeoJSON de la finca.")

    for etiqueta, valor in (
        ("altura_cm_1", body.altura_cm_1),
        ("altura_cm_2", body.altura_cm_2),
        ("altura_cm_3", body.altura_cm_3),
    ):
        if not (0 < valor <= 200):
            raise HTTPException(
                400,
                f"La medida '{etiqueta}' debe ser un numero mayor que 0 y menor o igual a 200 cm "
                f"(se recibio {valor}).",
            )

    if body.fecha:
        try:
            datetime.strptime(body.fecha, "%Y-%m-%d")
        except ValueError:
            raise HTTPException(400, f"La fecha '{body.fecha}' debe tener el formato AAAA-MM-DD.")
        fecha = body.fecha
    else:
        fecha = date.today().isoformat()

    guardar_fila_aforo(nombre, fecha, body.altura_cm_1, body.altura_cm_2, body.altura_cm_3)

    return {"mensaje": f"Medicion guardada para '{nombre}' el {fecha}.", "potrero": nombre, "fecha": fecha}


@app.post("/api/predecir")
def predecir():
    if not GEOJSON_PATH.exists():
        raise HTTPException(
            400,
            "No se encontro el GeoJSON de potreros en "
            f"'{GEOJSON_PATH}'. Exportalo desde SAD primero.",
        )

    ee_project = os.environ.get("EE_PROJECT")
    if not ee_project:
        raise HTTPException(
            400,
            "Falta configurar el proyecto de Google Cloud para Earth Engine: "
            "define la variable de entorno EE_PROJECT en backend/.env.",
        )

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    fecha_hoy = date.today().isoformat()

    # Asegura que potreros_estado.xlsx ya exista con el nombre de la finca
    # (FINCA_NOMBRE) ANTES de invocar el modelo, para que este lo herede
    # en vez de caer en su propio valor por defecto generico.
    cargar_estado()

    try:
        resultado_indices = subprocess.run(
            [
                sys.executable,
                str(EARTH_ENGINE_SCRIPT),
                "--geojson", str(GEOJSON_PATH),
                "--fecha", fecha_hoy,
                "--proyecto", ee_project,
                "--salida", str(INDICES_TEMP_PATH),
            ],
            capture_output=True,
            text=True,
            cwd=str(EARTH_ENGINE_SCRIPT.parent),
        )
    except OSError as error:
        raise HTTPException(500, f"No se pudo ejecutar el script de Earth Engine: {error}")

    if resultado_indices.returncode != 0:
        detalle = ultimas_lineas(resultado_indices.stdout + "\n" + resultado_indices.stderr)
        raise HTTPException(502, f"No se pudieron obtener los indices satelitales:\n{detalle}")

    agregar_indices_al_historial(INDICES_TEMP_PATH, fecha_hoy)

    if not MODELO_SCRIPT.exists():
        raise HTTPException(
            501,
            "El modelo de prediccion de biomasa todavia no esta implementado "
            "(se construira cuando haya datos de aforo suficientes para entrenarlo). "
            f"Los indices satelitales de hoy si se calcularon correctamente y quedaron en '{INDICES_HISTORIAL_PATH}'.",
        )

    argumentos_modelo = [
        sys.executable,
        str(MODELO_SCRIPT),
        "--geojson", str(GEOJSON_PATH),
        "--indices", str(INDICES_HISTORIAL_PATH),
        "--aforo", str(AFORO_PATH),
        "--salida", str(ESTADO_PATH),
        "--modelo-guardado", str(MODELO_JOBLIB_PATH),
        "--umbral-rojo", str(UMBRAL_ROJO),
        "--umbral-ambar", str(UMBRAL_AMBAR),
        "--modo", "auto",
    ]
    if CLIMA_PATH:
        if not CLIMA_PATH.exists():
            raise HTTPException(
                400,
                f"CLIMA_PATH apunta a '{CLIMA_PATH}', pero ese archivo no existe. "
                "Corrige la ruta en backend/.env o quita esa variable si todavia no tienes datos de clima.",
            )
        argumentos_modelo += ["--clima", str(CLIMA_PATH)]

    try:
        resultado_modelo = subprocess.run(
            argumentos_modelo,
            capture_output=True,
            text=True,
            cwd=str(MODELO_SCRIPT.parent),
        )
    except OSError as error:
        raise HTTPException(500, f"No se pudo ejecutar el modelo de prediccion: {error}")

    if resultado_modelo.returncode != 0:
        detalle = ultimas_lineas(resultado_modelo.stdout + "\n" + resultado_modelo.stderr)
        raise HTTPException(502, f"No se pudo generar la prediccion de biomasa:\n{detalle}")

    return cargar_estado()
